"""All reading and writing of the Excel workbook.

This is the ONLY module that imports openpyxl. Everything else in the
application works with plain Python dictionaries, where each dict maps a
column name to its value:

    {"Company ID": "C001", "Company Name": "Safaricom PLC", ...}

Two consequences worth understanding:

1. The AI and scoring modules can be tested with plain dicts, with no
   Excel file anywhere in sight.
2. If this project ever outgrows Excel, only this file changes.
"""

import logging
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app import schema

logger = logging.getLogger(__name__)


class ExcelError(Exception):
    """Raised when the workbook is missing, malformed, or cannot be written."""


# --- Creating the workbook -------------------------------------------------

# Column widths that make the sheet readable without manual resizing.
# Anything not listed gets DEFAULT_WIDTH.
DEFAULT_WIDTH = 18
COLUMN_WIDTHS = {
    "Company Name": 28,
    "Website": 30,
    "What They Do": 45,
    "Technology Function": 35,
    "Internship Evidence": 45,
    "Qualification Reason": 50,
    "Careers URL": 30,
    "Internship URL": 30,
    "Relevant Roles": 30,
    "Email": 30,
    "LinkedIn": 30,
    "Why This Contact": 40,
    "Subject": 40,
    "Email Body": 70,
    "Reply Summary": 40,
    "Next Action": 25,
    "Notes": 30,
    "Value": 60,
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")   # dark blue
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)


def _style_sheet(sheet: Worksheet, columns: list[str]) -> None:
    """Write the header row, style it, freeze it, and set column widths.

    Freezing the header means the column names stay visible while you
    scroll -- important when you are reviewing email drafts by hand.
    """
    sheet.append(columns)

    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left")

        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = COLUMN_WIDTHS.get(name, DEFAULT_WIDTH)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def create_workbook(path: Path, overwrite: bool = False) -> Path:
    """Create a fresh workbook with all four sheets and default settings.

    Args:
        path: Where to write the .xlsx file.
        overwrite: If False (the default) and the file already exists,
                   raise rather than destroy existing data. Refusing to
                   silently overwrite your database is deliberate.

    Returns:
        The path that was written.
    """
    if path.exists() and not overwrite:
        raise ExcelError(
            f"Workbook already exists at {path}. "
            "Use --overwrite if you really want to replace it "
            "(this deletes all existing data)."
        )

    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    # A new Workbook comes with one default sheet; rename it rather than
    # creating a fifth and deleting one.
    first = workbook.active
    first.title = schema.COMPANIES
    _style_sheet(first, schema.COMPANY_COLUMNS)

    for name in [schema.CONTACTS, schema.OUTREACH, schema.SETTINGS]:
        sheet = workbook.create_sheet(name)
        _style_sheet(sheet, schema.SHEET_COLUMNS[name])

    # Pre-fill the Settings sheet with the profile.
    settings_sheet = workbook[schema.SETTINGS]
    for key, value in schema.DEFAULT_SETTINGS:
        settings_sheet.append([key, value])

    _save(workbook, path)
    logger.info("Created workbook at %s with sheets: %s",
                path, ", ".join(schema.SHEET_NAMES))
    return path


# --- Low-level open/save ---------------------------------------------------

def _open(path: Path) -> Workbook:
    """Open the workbook, with clear errors for the two common failures."""
    if not path.exists():
        raise ExcelError(
            f"Workbook not found at {path}. "
            "Run `python main.py init-excel` to create it."
        )
    try:
        return load_workbook(path)
    except PermissionError as exc:
        raise ExcelError(
            f"Cannot open {path} -- it is probably open in Excel. "
            "Close it and try again."
        ) from exc
    except Exception as exc:
        raise ExcelError(f"Could not read {path}: {exc}") from exc


def _save(workbook: Workbook, path: Path) -> None:
    """Save the workbook, with a clear error if the file is locked.

    The PermissionError case happens constantly in real use: you left the
    workbook open in Excel and Python cannot write to it. A plain traceback
    is confusing, so we say exactly what to do.
    """
    try:
        workbook.save(path)
    except PermissionError as exc:
        raise ExcelError(
            f"Cannot write to {path} -- the file is open in Excel. "
            "Close it and run the command again."
        ) from exc
    except Exception as exc:
        raise ExcelError(f"Could not save {path}: {exc}") from exc


def _get_sheet(workbook: Workbook, name: str) -> Worksheet:
    """Fetch a sheet by name, listing what does exist if it is missing."""
    if name not in workbook.sheetnames:
        raise ExcelError(
            f"Sheet '{name}' is missing from the workbook. "
            f"Found: {', '.join(workbook.sheetnames)}. "
            "The workbook may be corrupted -- recreate it with "
            "`python main.py init-excel --overwrite`."
        )
    return workbook[name]


def _cell_value(value: Any) -> Any:
    """Normalise a cell value for use in Python.

    openpyxl returns None for empty cells; we convert to "" so downstream
    code can do `.strip()` without a None check everywhere. Dates are
    converted to ISO strings so they compare and serialise predictably.
    """
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()[:10]
    return value


# --- Reading ---------------------------------------------------------------

def read_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    """Read a sheet into a list of dicts, one per row.

    The header row defines the keys. Completely blank rows are skipped --
    they are common when you delete a row by hand in Excel.

    Adds a private "_row" key holding the real spreadsheet row number, so
    update functions know exactly which row to write back to.
    """
    workbook = _open(path)
    try:
        sheet = _get_sheet(workbook, sheet_name)
        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            return []

        headers = [_cell_value(h) for h in rows[0]]

        records: list[dict[str, Any]] = []
        for row_number, row in enumerate(rows[1:], start=2):
            values = [_cell_value(v) for v in row]
            # Skip rows where every cell is empty.
            if not any(str(v).strip() for v in values):
                continue

            record: dict[str, Any] = {}
            for header, value in zip(headers, values):
                if header:  # ignore unnamed trailing columns
                    record[str(header)] = value
            record["_row"] = row_number
            records.append(record)

        logger.debug("Read %d rows from '%s'", len(records), sheet_name)
        return records
    finally:
        workbook.close()


def read_companies(path: Path) -> list[dict[str, Any]]:
    """Read all companies."""
    return read_sheet(path, schema.COMPANIES)


def read_contacts(path: Path) -> list[dict[str, Any]]:
    """Read all contacts."""
    return read_sheet(path, schema.CONTACTS)


def read_outreach(path: Path) -> list[dict[str, Any]]:
    """Read all outreach records."""
    return read_sheet(path, schema.OUTREACH)


def read_settings(path: Path) -> dict[str, str]:
    """Read the Settings sheet as a plain key -> value dictionary."""
    rows = read_sheet(path, schema.SETTINGS)
    return {
        str(row.get("Key", "")).strip(): str(row.get("Value", "")).strip()
        for row in rows
        if str(row.get("Key", "")).strip()
    }


# --- Writing ---------------------------------------------------------------

def backup(path: Path, tag: str = "") -> Path | None:
    """Copy the workbook to data/backups/ before a bulk change.

    Cheap insurance. The workbook is the only copy of your research, and
    a bulk operation that goes wrong should never be unrecoverable.
    """
    import shutil
    from datetime import datetime

    if not path.exists():
        return None

    backup_dir = path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    suffix = f"-{tag}" if tag else ""
    destination = backup_dir / f"{path.stem}-{stamp}{suffix}.xlsx"

    try:
        shutil.copy2(path, destination)
    except OSError as exc:
        logger.warning("Could not write a backup: %s", exc)
        return None

    # Keep the 20 most recent; older ones are rarely useful.
    backups = sorted(backup_dir.glob(f"{path.stem}-*.xlsx"))
    for old in backups[:-20]:
        try:
            old.unlink()
        except OSError:
            pass

    logger.debug("Backed up the workbook to %s", destination)
    return destination


def append_rows(path: Path, sheet_name: str, records: list[dict[str, Any]]) -> int:
    """Append records to a sheet, matching dict keys to column headers.

    Keys that do not match a column are ignored (with a warning, so a typo
    does not vanish silently). Missing keys become empty cells.

    Returns the number of rows appended.
    """
    if not records:
        logger.info("Nothing to append to '%s'.", sheet_name)
        return 0

    workbook = _open(path)
    try:
        sheet = _get_sheet(workbook, sheet_name)
        columns = schema.SHEET_COLUMNS.get(sheet_name)
        if columns is None:
            raise ExcelError(f"No column definition for sheet '{sheet_name}'.")

        known = set(columns)
        for record in records:
            unknown_keys = set(record) - known - {"_row"}
            if unknown_keys:
                logger.warning(
                    "Ignoring unknown column(s) for '%s': %s",
                    sheet_name, ", ".join(sorted(unknown_keys)),
                )
            sheet.append([record.get(column, "") for column in columns])

        _save(workbook, path)
        logger.info("Appended %d row(s) to '%s'.", len(records), sheet_name)
        return len(records)
    finally:
        workbook.close()


def update_row(
    path: Path, sheet_name: str, row_number: int, updates: dict[str, Any]
) -> None:
    """Update specific cells in one existing row.

    Args:
        row_number: The spreadsheet row (the "_row" value from a read).
        updates: column name -> new value. Columns not mentioned are
                 left untouched, so this never clobbers your manual edits.
    """
    workbook = _open(path)
    try:
        sheet = _get_sheet(workbook, sheet_name)
        headers = [_cell_value(c.value) for c in sheet[1]]
        index_of = {str(name): i + 1 for i, name in enumerate(headers) if name}

        if row_number < 2 or row_number > sheet.max_row:
            raise ExcelError(
                f"Row {row_number} is out of range for '{sheet_name}' "
                f"(rows 2..{sheet.max_row})."
            )

        for column, value in updates.items():
            if column not in index_of:
                logger.warning(
                    "Column '%s' does not exist in '%s' -- skipping.",
                    column, sheet_name,
                )
                continue
            sheet.cell(row=row_number, column=index_of[column], value=value)

        _save(workbook, path)
        logger.debug("Updated row %d of '%s'.", row_number, sheet_name)
    finally:
        workbook.close()


def update_rows(
    path: Path, sheet_name: str, updates_by_row: dict[int, dict[str, Any]]
) -> int:
    """Update many rows in a single open/save cycle.

    Opening and saving an .xlsx file is slow. When Stage 3 qualifies ten
    companies we want one save, not ten, so this batches them.

    IMPORTANT: this opens the workbook NOW and saves immediately. Never
    hold an open Workbook across slow work (web fetches, AI calls) and
    save it afterwards -- anything written to the file in the meantime
    would be silently overwritten. That bug cost 32 rows once; the fix
    is to gather your changes in a plain dict and call this at the end.

    Returns the number of rows updated.
    """
    if not updates_by_row:
        return 0

    workbook = _open(path)
    try:
        sheet = _get_sheet(workbook, sheet_name)
        headers = [_cell_value(c.value) for c in sheet[1]]
        index_of = {str(name): i + 1 for i, name in enumerate(headers) if name}

        updated = 0
        for row_number, updates in updates_by_row.items():
            if row_number < 2 or row_number > sheet.max_row:
                logger.warning(
                    "Skipping out-of-range row %d in '%s'.", row_number, sheet_name
                )
                continue
            for column, value in updates.items():
                if column not in index_of:
                    logger.warning(
                        "Column '%s' does not exist in '%s' -- skipping.",
                        column, sheet_name,
                    )
                    continue
                sheet.cell(row=row_number, column=index_of[column], value=value)
            updated += 1

        _save(workbook, path)
        logger.info("Updated %d row(s) in '%s'.", updated, sheet_name)
        return updated
    finally:
        workbook.close()


# --- ID generation ---------------------------------------------------------

def next_id(existing: list[dict[str, Any]], id_column: str, prefix: str) -> str:
    """Generate the next sequential ID, e.g. C001, C002.

    Reads the highest existing numeric suffix and adds one, so IDs stay
    unique even if you deleted rows in the middle. Non-numeric or
    hand-edited IDs are ignored rather than crashing.
    """
    highest = 0
    for record in existing:
        raw = str(record.get(id_column, "")).strip()
        if raw.upper().startswith(prefix.upper()):
            suffix = raw[len(prefix):]
            if suffix.isdigit():
                highest = max(highest, int(suffix))
    return f"{prefix}{highest + 1:03d}"


def today() -> str:
    """Today's date as an ISO string, for the Date * columns."""
    return date.today().isoformat()


# --- Health check ----------------------------------------------------------

def validate_workbook(path: Path) -> list[str]:
    """Check the workbook has the expected sheets and columns.

    Returns a list of human-readable problems. An empty list means healthy.
    Used by `python main.py check-excel` so you can diagnose a workbook you
    have been editing by hand.
    """
    problems: list[str] = []

    try:
        workbook = _open(path)
    except ExcelError as exc:
        return [str(exc)]

    try:
        for sheet_name in schema.SHEET_NAMES:
            if sheet_name not in workbook.sheetnames:
                problems.append(f"Missing sheet: '{sheet_name}'")
                continue

            sheet = workbook[sheet_name]
            headers = [str(_cell_value(c.value)) for c in sheet[1]]
            expected = schema.SHEET_COLUMNS[sheet_name]

            for column in expected:
                if column not in headers:
                    problems.append(
                        f"Sheet '{sheet_name}' is missing column: '{column}'"
                    )
    finally:
        workbook.close()

    return problems
