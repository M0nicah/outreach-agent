"""Stage 2 tests: the workbook round-trips data correctly.

Each test builds a throwaway workbook in a temp directory, so these never
touch your real data/internship_outreach.xlsx.

Run with:   python -m unittest discover tests
"""

import tempfile
import unittest
from pathlib import Path

from app import schema
from app.excel import (
    ExcelError,
    append_rows,
    create_workbook,
    next_id,
    read_companies,
    read_settings,
    update_row,
    update_rows,
    validate_workbook,
)


class ExcelTestCase(unittest.TestCase):
    """Base class that gives each test a fresh workbook."""

    def setUp(self):
        self._tempdir = tempfile.TemporaryDirectory()
        self.path = Path(self._tempdir.name) / "test.xlsx"
        create_workbook(self.path)

    def tearDown(self):
        self._tempdir.cleanup()


class TestCreateWorkbook(ExcelTestCase):
    def test_all_four_sheets_exist_with_correct_columns(self):
        self.assertEqual(validate_workbook(self.path), [])

    def test_settings_are_prefilled_with_profile(self):
        settings = read_settings(self.path)
        self.assertEqual(settings["Name"], "Monica Wughanga Masae")
        self.assertEqual(settings["University"], "Open University of Kenya")
        self.assertEqual(settings["Degree"], "BSc Data Science")

    def test_refuses_to_overwrite_without_permission(self):
        with self.assertRaises(ExcelError):
            create_workbook(self.path)

    def test_overwrite_works_when_explicitly_requested(self):
        append_rows(self.path, schema.COMPANIES, [{"Company Name": "Test Co"}])
        self.assertEqual(len(read_companies(self.path)), 1)

        create_workbook(self.path, overwrite=True)
        self.assertEqual(len(read_companies(self.path)), 0)


class TestReadWrite(ExcelTestCase):
    def test_append_then_read_returns_the_same_values(self):
        append_rows(
            self.path,
            schema.COMPANIES,
            [
                {"Company ID": "C001", "Company Name": "Alpha Ltd", "Country": "Kenya"},
                {"Company ID": "C002", "Company Name": "Beta Ltd", "Country": "Kenya"},
            ],
        )
        companies = read_companies(self.path)

        self.assertEqual(len(companies), 2)
        self.assertEqual(companies[0]["Company Name"], "Alpha Ltd")
        self.assertEqual(companies[1]["Company ID"], "C002")

    def test_missing_keys_become_empty_not_errors(self):
        append_rows(self.path, schema.COMPANIES, [{"Company Name": "Sparse Ltd"}])
        company = read_companies(self.path)[0]
        self.assertEqual(company["Website"], "")

    def test_row_numbers_let_us_update_the_right_row(self):
        append_rows(
            self.path,
            schema.COMPANIES,
            [{"Company Name": "First"}, {"Company Name": "Second"}],
        )
        companies = read_companies(self.path)
        # Row 1 is the header, so data starts at row 2.
        self.assertEqual(companies[0]["_row"], 2)
        self.assertEqual(companies[1]["_row"], 3)

    def test_update_changes_only_the_named_columns(self):
        append_rows(
            self.path,
            schema.COMPANIES,
            [{"Company Name": "Gamma Ltd", "Country": "Kenya"}],
        )
        row = read_companies(self.path)[0]["_row"]

        update_row(
            self.path,
            schema.COMPANIES,
            row,
            {"Research Status": schema.RESEARCH_QUALIFY, "Total Score": 88},
        )

        company = read_companies(self.path)[0]
        self.assertEqual(company["Research Status"], "QUALIFY")
        self.assertEqual(company["Total Score"], 88)
        # Untouched columns must survive -- this protects your manual edits.
        self.assertEqual(company["Company Name"], "Gamma Ltd")
        self.assertEqual(company["Country"], "Kenya")

    def test_batch_update_writes_every_row(self):
        append_rows(
            self.path,
            schema.COMPANIES,
            [{"Company Name": f"Co {i}"} for i in range(3)],
        )
        updated = update_rows(
            self.path,
            schema.COMPANIES,
            {
                2: {"Research Status": schema.RESEARCH_QUALIFY},
                3: {"Research Status": schema.RESEARCH_REJECT},
                4: {"Research Status": schema.RESEARCH_NEEDS_REVIEW},
            },
        )
        self.assertEqual(updated, 3)

        statuses = [c["Research Status"] for c in read_companies(self.path)]
        self.assertEqual(statuses, ["QUALIFY", "REJECT", "NEEDS_REVIEW"])

    def test_unknown_column_is_ignored_not_crashed_on(self):
        append_rows(
            self.path,
            schema.COMPANIES,
            [{"Company Name": "Delta Ltd", "Not A Real Column": "x"}],
        )
        self.assertEqual(read_companies(self.path)[0]["Company Name"], "Delta Ltd")

    def test_out_of_range_row_raises_clear_error(self):
        with self.assertRaises(ExcelError):
            update_row(self.path, schema.COMPANIES, 999, {"Company Name": "Nope"})


class TestErrorHandling(unittest.TestCase):
    def test_missing_file_gives_actionable_message(self):
        missing = Path(tempfile.gettempdir()) / "definitely-not-here-12345.xlsx"
        with self.assertRaises(ExcelError) as ctx:
            read_companies(missing)
        self.assertIn("init-excel", str(ctx.exception))

    def test_validate_reports_a_missing_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "broken.xlsx"
            create_workbook(path)

            # Delete a sheet the way a user might, by hand.
            from openpyxl import load_workbook

            workbook = load_workbook(path)
            del workbook[schema.CONTACTS]
            workbook.save(path)

            problems = validate_workbook(path)
            self.assertTrue(any("Contacts" in p for p in problems), problems)


class TestNextId(unittest.TestCase):
    def test_starts_at_one_when_empty(self):
        self.assertEqual(next_id([], "Company ID", "C"), "C001")

    def test_continues_from_the_highest_existing(self):
        existing = [{"Company ID": "C001"}, {"Company ID": "C007"}]
        self.assertEqual(next_id(existing, "Company ID", "C"), "C008")

    def test_ignores_hand_edited_non_numeric_ids(self):
        existing = [{"Company ID": "C002"}, {"Company ID": "C-oops"}]
        self.assertEqual(next_id(existing, "Company ID", "C"), "C003")


if __name__ == "__main__":
    unittest.main()
